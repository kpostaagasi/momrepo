import os, pickle, pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
r=pickle.load(open("res.pkl","rb")); U=r["U"]; DATE=r["DATE"].strftime("%Y-%m-%d")
A=pd.concat(U.values()).sort_values("MOMADJ",ascending=False)
K=["1a","3a","6a","12a"]
COLS=[("Enstrüman","Enstrüman",None),("Evren","Evren",None),("Son","Son","#,##0.00")]
for p,lab,f in [("ret","Getiri","0.0%"),("vol","Ort hacim","#,##0"),("vr","Hacim oranı","0.00"),("zr","z getiri","0.00"),("zv","z hacim","0.00"),("mom","MOM","0.00")]:
    COLS+=[(f"{p}_{k}",f"{lab} {k}",f) for k in K]
COLS+=[("ZRET","z getiri bileşik","0.00"),("ZVOL","z hacim bileşik","0.00"),("MOM","MOM bileşik","0.00"),("MOMADJ","MOM_ADJ bileşik","0.00"),("Kadran","Kadran",None)]
assert len(COLS)==32
wb=Workbook(); F=Font(name="Arial",size=10); HF=Font(name="Arial",bold=True,color="FFFFFF"); HFill=PatternFill("solid",fgColor="12314F")
ws=wb.active; ws.title="Yöntem"
nq4=int(((A.MOM>0)&(A.ZRET<0)).sum()); nm=int(A.MOM.notna().sum())
prox=[(v,U["Emtia"].loc[U["Emtia"].Not==v].Enstrüman.tolist()) for v in U["Emtia"].Not.unique() if v]
lines=[("Getiri–Hacim Momentum Çalışması",""),("Veri tarihi",DATE),("Kaynak","Yahoo Finance günlük (v8 chart, range=3y). Veri tarihi günü tamamlanmış son seanstır; bugünün kısmi barları atılır."),("",""),
("ADIMLAR",""),
("1. Veri filtresi","En az 505 bar (252 günlük getiri + 504 günlük hacim tabanı için). Son barı veri tarihinden 7 günden eski olan düşer."),
("2. Getiri","ret_k = P_t / P_(t-k) − 1; k = 21/63/126/252 (1a/3a/6a/12a). Düzeltilmiş kapanış."),
("2b. Getiri serileri (^TNX, ^TYX, ^IRX)","Fiyat-eşdeğer getiri: ret_k = −D × Δy_k / 100; D = 8.5 / 17.0 / 0.25. Böylece tahvil ETF'leriyle aynı yönde okunur."),
("3. Hacim oranı","volratio_k = ort(hacim, son k gün) / ort(hacim, son 504 gün). Sıfır hacim NaN sayılır."),
("4. Getiri z-skoru","z_ret_k = kesitsel_z(winsorize(ret_k, %2, %98)), ±3'te kırpılır"),
("5. Hacim standardizasyonu","z_vol_k = kesitsel_z(winsorize(ln(volratio_k), %2, %98)), ±3'te kırpılır; volratio ≤ 0 → NaN"),
("6. MOM","MOM_k = z_ret_k × z_vol_k"),
("7. MOMADJ","MOMADJ_k = z_ret_k × exp(clip(z_vol_k, −2, 2) / 2); çarpan 0.37x–2.72x. Hacimsiz satırda çarpan 1."),
("8. Bileşik","Ağırlıklar 1a 0.40 / 3a 0.30 / 6a 0.20 / 12a 0.10; MOM, MOMADJ, ZRET, ZVOL için aynı"),
("9. Sıralama","MOMADJ azalan. Teyit: ham MOM + kadran."),
("",""),("KADRANLAR",""),
("Q1 · Teyitli yükseliş","ZRET ≥ 0, ZVOL ≥ 0"),("Q2 · Teyitsiz yükseliş","ZRET ≥ 0, ZVOL < 0 (hacimsiz ralli)"),("Q3 · Teyitli düşüş","ZRET < 0, ZVOL ≥ 0 (dağıtım)"),("Q4 · İlgisiz düşüş","ZRET < 0, ZVOL < 0 (ham MOM burada yanıltıcı şekilde pozitif)"),("Hacimsiz","Hacim yok ya da güvenilmez; MOM hesaplanmaz, z_ret evrene dahil"),
("",""),("MOM İŞARET UYARISI",f"Bu koşuda hacimli {nm} satırın {nq4}'inde (%{nq4/nm*100:.0f}) getiri negatifken ham MOM pozitif. Sıralamada ham MOM kullanılmaz."),
("",""),("HACİM KALİTESİ (emtia, son 120 bar)","Geçme şartı: ≥60 gözlem ve kapsama ≥%60; medyan > 0; p90/medyan ≤ 5.0; medyanın %20'si altındaki gün payı ≤ %35"),
("Vekil eşleşmeleri","Altın→GLD, Gümüş→SLV, Platin→PPLT, Paladyum→PALL (yalnızca test başarısızsa; tarihler normalize edilerek birleştirilir)"),
]
for v,names in prox: lines.append((", ".join(names),v))
lines+=[("Getiri serileri","^TNX, ^TYX, ^IRX hacimsizdir (vekil bağlanmaz)"),("",""),("EVREN NOTLARI",""),
("BIST 30","Borsa İstanbul 1 Tem–30 Eyl 2026 dönemi 30 pay (TRALT dahil)"),
("BIST 100",f"Resmi liste değil: {r['nbroad']} paylık geniş evrenden 60 günlük ort. TL işlem hacmine göre ilk 100"),
("Nasdaq 100","Bileşen listesi: api.nasdaq.com (nasdaq100), çalışma günü"),
("Alt küme notu","BIST 30, BIST 100'ün alt kümesidir; z-skorlar evren içinde hesaplandığı için aynı hissenin iki sekmedeki skoru farklıdır."),
("Evrenler arası","Skorlar evrenler arasında karşılaştırılamaz; TUMU sekmesi yalnızca filtreleme içindir."),
("Olay işareti","1a getiri ≥ %30 ve 12a getiri < 0 olan isimler Enstrüman hücresinde yorumla işaretlidir."),
("Düşen semboller","; ".join(f"{s} ({why})" for s,why in r["dropped"])),
("Önceki koşu", (f"Karşılaştırma tabanı: {r['prev']['tarih']}; kadran değiştiren {len(r['prev']['changes'])} satır (PDF sayfa 1)") if r["prev"] else "Önceki koşu kaydı yok; karşılaştırma yapılmadı."),
]
for i,(a,b) in enumerate(lines,1):
    ws.cell(i,1,a).font=Font(name="Arial",bold=a.isupper() or i==1,size=12 if i==1 else 10,color="12314F" if a.isupper() or i==1 else "000000")
    c=ws.cell(i,2,b); c.font=F; c.alignment=Alignment(wrap_text=True,vertical="top")
    ws.cell(i,1).alignment=Alignment(wrap_text=True,vertical="top")
ws.column_dimensions["A"].width=34; ws.column_dimensions["B"].width=110
def sheet(name,t):
    s=wb.create_sheet(name)
    for j,(k,lab,f) in enumerate(COLS,1):
        c=s.cell(1,j,lab); c.font=HF; c.fill=HFill; c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
        s.column_dimensions[get_column_letter(j)].width=24 if j==1 else (16 if j in (2,32) else (13 if f=="#,##0" else 10))
    s.column_dimensions["AF"].width=22
    for i,(_,row) in enumerate(t.iterrows(),2):
        for j,(k,lab,f) in enumerate(COLS,1):
            v=row[k]
            if isinstance(v,(float,np.floating)) and not np.isfinite(v): v=None
            elif isinstance(v,np.floating): v=float(v)
            c=s.cell(i,j,v); c.font=F
            if f: c.number_format=f
        note=[x for x in [row.Not, "Olay işareti: 1a ≥ %30, 12a < 0; haber akışı kontrol edilmeli" if row.Olay else ""] if x]
        if note: s.cell(i,1).comment=Comment("\n".join(note),"Model")
    s.freeze_panes="D2"; s.auto_filter.ref=f"A1:{get_column_letter(32)}{len(t)+1}"; s.row_dimensions[1].height=30
for n,key in [("Emtia","Emtia"),("Tahvil","Tahvil / faiz"),("Nasdaq 100","Nasdaq 100"),("BIST 100","BIST 100"),("BIST 30","BIST 30")]: sheet(n,U[key])
sheet("TUMU",A)
out=os.environ.get("OUT_DIR","out")+f"/Momentum_Calismasi_{DATE}.xlsx"; wb.save(out); print(out, wb.sheetnames)
